"""Import vehicles from Avito autoload Excel (.xlsx)."""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import Path
from typing import BinaryIO, Iterable

from django.core.files.uploadedfile import SimpleUploadedFile
from PIL import Image

from utils.html_sanitize import sanitize_html
from utils.safe_http import fetch_url_bytes, is_safe_request_url

from .avito import parse_avito_item_id
from .engine_type import detect_engine_type
from .listing_ingest import (
    MAX_GALLERY_UPLOADS,
    attach_vehicle_images,
    create_vehicle_draft,
    get_or_create_brand,
    get_or_create_category,
)
from .models import Vehicle

logger = logging.getLogger(__name__)

SHEET_NAME = "Автомобили-С пробегом"
MAX_IMAGE_BYTES = 15 * 1024 * 1024
IMAGE_HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
    ),
}

_PUNCT_RE = re.compile(r"[^\w\s]+", re.UNICODE)
_SPACE_RE = re.compile(r"\s+")
_DIGITS_RE = re.compile(r"\d+")

SPEC_KEYS = (
    ("GenerationId", "generation"),
    ("ModificationId", "modification"),
    ("ComplectationId", "complectation"),
    ("VIN", "vin"),
    ("Owners", "owners"),
    ("PTS", "pts"),
    ("DriveType", "drive"),
    ("EngineSize", "engine_size"),
    ("Doors", "doors"),
    ("WheelType", "wheel"),
    ("Accident", "accident"),
    ("FuelType", "fuel"),
)


@dataclass
class AvitoListing:
    row_number: int
    avito_id: int | None
    make: str
    model: str
    title: str
    short_title: str
    year: int
    mileage: int
    color: str
    price_rub: Decimal | None
    description: str
    body_type: str
    transmission: str
    fuel_type: str
    horsepower: int | None
    image_urls: list[str] = field(default_factory=list)
    specs: dict = field(default_factory=dict)


@dataclass
class ImportReport:
    created: int = 0
    skipped: int = 0
    linked: int = 0
    errors: list[str] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    created_ids: list[int] = field(default_factory=list)
    linked_ids: list[int] = field(default_factory=list)
    skipped_reasons: list[str] = field(default_factory=list)
    linked_reasons: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return (
            f"Создано: {self.created}, привязан Avito ID: {self.linked}, "
            f"пропущено: {self.skipped}, "
            f"ошибок: {len(self.errors)}, предупреждений: {len(self.warnings)}."
        )


def normalize_text(value: str | None) -> str:
    text = (value or "").replace("\u00a0", " ").replace("\u202f", " ")
    text = text.casefold()
    text = _PUNCT_RE.sub(" ", text)
    text = _SPACE_RE.sub(" ", text).strip()
    return text


def _cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _parse_int(value, default: int = 0) -> int:
    if value is None or value == "":
        return default
    if isinstance(value, bool):
        return default
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)
    digits = "".join(_DIGITS_RE.findall(str(value).replace("\u00a0", " ")))
    if not digits:
        return default
    try:
        return int(digits)
    except ValueError:
        return default


def _parse_price(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        try:
            return Decimal(str(int(value)))
        except (InvalidOperation, ValueError, OverflowError):
            return None
    digits = "".join(_DIGITS_RE.findall(str(value).replace("\u00a0", " ")))
    if not digits:
        return None
    try:
        return Decimal(digits)
    except InvalidOperation:
        return None


def _split_image_urls(raw: str | None) -> list[str]:
    if not raw:
        return []
    parts = re.split(r"\s*\|\s*", str(raw).strip())
    urls: list[str] = []
    for part in parts:
        url = part.strip()
        if not url:
            continue
        if url.startswith("http://"):
            url = "https://" + url[len("http://") :]
        urls.append(url)
    return urls[:MAX_GALLERY_UPLOADS]


def _find_header_row(rows: list[tuple]) -> tuple[int, list[str]]:
    for idx, row in enumerate(rows[:12]):
        cells = [_cell_str(c) for c in row]
        if "Id" in cells and "AvitoId" in cells and "Make" in cells:
            return idx, cells
    raise ValueError(
        f"Не найдены заголовки Id/AvitoId/Make на листе «{SHEET_NAME}»."
    )


def parse_avito_xlsx(source: str | Path | BinaryIO) -> list[AvitoListing]:
    """Parse Avito autoload workbook into listing rows."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError(
            "Нужен пакет openpyxl. Установите: pip install openpyxl"
        ) from exc

    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if SHEET_NAME not in wb.sheetnames:
            raise ValueError(
                f"В файле нет листа «{SHEET_NAME}». Есть: {', '.join(wb.sheetnames)}"
            )
        ws = wb[SHEET_NAME]
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    if not rows:
        return []

    header_idx, headers = _find_header_row(rows)
    col = {name: i for i, name in enumerate(headers) if name}

    def get(row, key: str):
        idx = col.get(key)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    listings: list[AvitoListing] = []
    for row_offset, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        make = _cell_str(get(row, "Make"))
        if not make:
            continue
        # Skip Avito instruction / validation helper rows.
        if make in {"Обязательный", "Подробнее о параметре", "Марка"}:
            continue
        model = _cell_str(get(row, "Model"))
        year = _parse_int(get(row, "Year"))
        if year < 1980:
            continue

        avito_raw = get(row, "AvitoId") or get(row, "Id")
        avito_id = parse_avito_item_id(avito_raw)
        title = _cell_str(get(row, "Title"))
        short_title = f"{make} {model}".strip()
        if not title:
            title = short_title

        specs: dict[str, str] = {}
        for excel_key, spec_key in SPEC_KEYS:
            val = _cell_str(get(row, excel_key))
            if val:
                specs[spec_key] = val
        color = _cell_str(get(row, "Color"))
        if color:
            specs["color"] = color

        listings.append(
            AvitoListing(
                row_number=row_offset,
                avito_id=avito_id,
                make=make,
                model=model,
                title=title,
                short_title=short_title,
                year=year,
                mileage=_parse_int(get(row, "Kilometrage")),
                color=color,
                price_rub=_parse_price(get(row, "Price")),
                description=_cell_str(get(row, "Description")),
                body_type=_cell_str(get(row, "BodyType")),
                transmission=_cell_str(get(row, "Transmission")),
                fuel_type=_cell_str(get(row, "FuelType")),
                horsepower=_parse_int(get(row, "Power")) or None,
                image_urls=_split_image_urls(_cell_str(get(row, "ImageUrls"))),
                specs=specs,
            )
        )
    return listings


def titles_match(listing: AvitoListing, vehicle: Vehicle) -> bool:
    site_title = normalize_text(vehicle.title)
    brand_name = normalize_text(getattr(vehicle.brand, "name", "") or "")
    site_brand_model = normalize_text(
        f"{brand_name} {normalize_text(vehicle.model)}".strip()
    )
    avito_title = normalize_text(listing.title)
    avito_short = normalize_text(listing.short_title)
    if not site_title:
        return False
    if site_title == avito_title:
        return True
    if site_brand_model and site_brand_model == avito_short:
        return True
    if site_title == avito_short:
        return True
    return False


def fingerprint_match(listing: AvitoListing, vehicle: Vehicle) -> bool:
    if int(vehicle.year or 0) != int(listing.year or 0):
        return False
    if int(vehicle.mileage or 0) != int(listing.mileage or 0):
        return False
    if normalize_text(vehicle.color) != normalize_text(listing.color):
        return False
    return titles_match(listing, vehicle)


def find_existing_vehicle(
    listing: AvitoListing,
    *,
    by_avito_id: dict[int, Vehicle],
    candidates: Iterable[Vehicle],
) -> Vehicle | None:
    if listing.avito_id and listing.avito_id in by_avito_id:
        return by_avito_id[listing.avito_id]
    for vehicle in candidates:
        if fingerprint_match(listing, vehicle):
            return vehicle
    return None


def _download_images(urls: list[str]) -> tuple[list[SimpleUploadedFile], list[str]]:
    uploads: list[SimpleUploadedFile] = []
    warnings: list[str] = []
    for index, url in enumerate(urls, start=1):
        if not is_safe_request_url(url):
            warnings.append(f"Небезопасный URL фото пропущен: {url[:80]}")
            continue
        try:
            raw, _final = fetch_url_bytes(
                url,
                max_bytes=MAX_IMAGE_BYTES,
                timeout=20,
                headers=IMAGE_HEADERS,
            )
            img = Image.open(BytesIO(raw))
            img.verify()
            img = Image.open(BytesIO(raw))
            fmt = (img.format or "JPEG").upper()
            ext_map = {"JPEG": ".jpg", "JPG": ".jpg", "PNG": ".png", "WEBP": ".webp"}
            ext = ext_map.get(fmt, ".jpg")
            uploads.append(
                SimpleUploadedFile(
                    f"avito_{index}{ext}",
                    raw,
                    content_type=f"image/{ext.lstrip('.')}",
                )
            )
        except Exception as exc:
            warnings.append(f"Фото не скачалось ({url[:60]}…): {exc}")
            logger.info("Avito image download failed url=%s err=%s", url[:120], exc)
    return uploads, warnings


def create_from_listing(
    listing: AvitoListing,
    *,
    download_photos: bool = True,
) -> tuple[Vehicle, list[str]]:
    brand, _ = get_or_create_brand(listing.make)
    category, _ = get_or_create_category(
        "Автомобили",
        body_type=listing.body_type,
        haystack=listing.description,
        mileage=listing.mileage,
    )
    vehicle = create_vehicle_draft(
        title=listing.short_title[:160] or listing.title[:160],
        brand=brand,
        category=category,
        model=listing.model,
        year=listing.year,
        mileage=listing.mileage,
        horsepower=listing.horsepower,
        transmission=listing.transmission,
        body_type=listing.body_type,
        color=listing.color,
        engine_type=detect_engine_type(listing.fuel_type),
        price_rub=listing.price_rub,
        description=sanitize_html(listing.description),
        specs=listing.specs,
        is_new=False,
    )
    warnings: list[str] = []
    update_fields: list[str] = []
    if listing.avito_id:
        vehicle.avito_item_id = listing.avito_id
        update_fields.append("avito_item_id")
    if update_fields:
        vehicle.save(update_fields=update_fields, skip_image_queue=True)

    if download_photos and listing.image_urls:
        uploads, photo_warnings = _download_images(listing.image_urls)
        warnings.extend(photo_warnings)
        if uploads:
            added, skipped = attach_vehicle_images(vehicle, uploads)
            if skipped:
                warnings.append(
                    f"Строка {listing.row_number}: пропущено файлов фото: {skipped}."
                )
            if not added:
                warnings.append(
                    f"Строка {listing.row_number}: фото не прикрепились."
                )
    return vehicle, warnings


def import_avito_listings(
    listings: list[AvitoListing],
    *,
    dry_run: bool = False,
    download_photos: bool = True,
) -> ImportReport:
    report = ImportReport()
    vehicles = list(
        Vehicle.objects.select_related("brand").only(
            "id",
            "title",
            "model",
            "year",
            "mileage",
            "color",
            "avito_item_id",
            "brand__name",
        )
    )
    by_avito_id = {
        int(v.avito_item_id): v for v in vehicles if v.avito_item_id
    }

    for listing in listings:
        try:
            existing = find_existing_vehicle(
                listing, by_avito_id=by_avito_id, candidates=vehicles
            )
            if existing is not None:
                # Fingerprint match without Avito ID → fill ID so site→Avito price sync works.
                if (
                    listing.avito_id
                    and not existing.avito_item_id
                    and listing.avito_id not in by_avito_id
                ):
                    reason = (
                        f"Строка {listing.row_number}: к «{existing.title}» "
                        f"(id={existing.pk}) привязан AvitoId {listing.avito_id}"
                    )
                    if dry_run:
                        report.linked += 1
                        report.linked_ids.append(existing.pk)
                        report.linked_reasons.append(reason + " (проверка)")
                        continue
                    existing.avito_item_id = listing.avito_id
                    existing.save(
                        update_fields=["avito_item_id"], skip_image_queue=True
                    )
                    by_avito_id[int(listing.avito_id)] = existing
                    report.linked += 1
                    report.linked_ids.append(existing.pk)
                    report.linked_reasons.append(reason)
                    continue

                report.skipped += 1
                reason = (
                    f"Строка {listing.row_number}: уже есть "
                    f"«{existing.title}» (id={existing.pk})"
                )
                if listing.avito_id and existing.avito_item_id == listing.avito_id:
                    reason += " — совпал AvitoId"
                elif existing.avito_item_id and listing.avito_id:
                    reason += (
                        f" — совпали название/год/пробег/цвет, "
                        f"на сайте уже другой AvitoId {existing.avito_item_id}"
                    )
                else:
                    reason += " — совпали название/год/пробег/цвет"
                report.skipped_reasons.append(reason)
                continue

            if dry_run:
                report.created += 1
                report.created_ids.append(0)
                continue

            vehicle, warnings = create_from_listing(
                listing, download_photos=download_photos
            )
            report.created += 1
            report.created_ids.append(vehicle.pk)
            report.warnings.extend(
                f"Строка {listing.row_number}: {w}" if not w.startswith("Строка") else w
                for w in warnings
            )
            vehicles.append(vehicle)
            if vehicle.avito_item_id:
                by_avito_id[int(vehicle.avito_item_id)] = vehicle
        except Exception as exc:
            logger.exception("Avito row import failed row=%s", listing.row_number)
            report.errors.append(f"Строка {listing.row_number}: {exc}")

    return report


def import_avito_xlsx(
    source: str | Path | BinaryIO,
    *,
    dry_run: bool = False,
    download_photos: bool = True,
) -> ImportReport:
    listings = parse_avito_xlsx(source)
    return import_avito_listings(
        listings, dry_run=dry_run, download_photos=download_photos
    )
